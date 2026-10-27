import tempfile
import unittest
from pathlib import Path
from zipfile import ZipFile

from lxml import etree

from openpyxl import Workbook
from pptx import Presentation
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

from survey_studio.smart_slides import (analyze_survey, default_analysis_plan, load_survey,
                                        render_analysis, validate_analysis_plan)
from survey_studio.smart_slides import BODY_FONT, BRAND_LOGO, HEADING_FONT, PALE


class SmartSlidesTests(unittest.TestCase):
    def test_real_csv_reproduces_reference_rating_and_subgroups(self):
        source = Path(__file__).parents[1] / "examples" / "data-set.csv"
        settings, headers, records, labels = load_survey(source)
        analysis = analyze_survey(headers, records, labels)
        rating = next(item for item in analysis["ratings"] if item["field"] == "d1")
        drive_cut = next(item for item in analysis["slides"].values()
                         if item["kind"] == "subgroup" and item["group_field"] == "a3_drive")
        groups = {item["name"]: item for item in drive_cut["groups"]}

        self.assertEqual(analysis["n"], 203)
        self.assertAlmostEqual(rating["summary"]["mean"], 8.0935960591, places=8)
        self.assertEqual(rating["summary"]["top_count"], 104)
        self.assertAlmostEqual(rating["summary"]["top_share"], 104 / 203)
        self.assertEqual(rating["summary"]["bottom_count"], 10)
        self.assertAlmostEqual(rating["summary"]["bottom_share"], 10 / 203)
        self.assertEqual({name: row["n"] for name, row in groups.items()}, {
            "BEV": 92, "Combustion engine + hybrid": 78, "PHEV": 33,
        })
        self.assertAlmostEqual(groups["BEV"]["mean"], 8.3804347826, places=8)
        self.assertFalse(analysis["coverage"]["demographics"])
        self.assertFalse(analysis["coverage"]["purchase_intent_or_rejection"])
        self.assertFalse(analysis["coverage"]["pricing_expectations"])
        self.assertFalse(analysis["coverage"]["emerging_brand_acceptance"])

    def test_real_data_renders_editable_charts_and_subgroup_table(self):
        source = Path(__file__).parents[1] / "examples" / "data-set.csv"
        settings, headers, records, labels = load_survey(source)
        analysis = analyze_survey(headers, records, labels)
        plan = default_analysis_plan(analysis)
        with tempfile.TemporaryDirectory() as folder:
            output_path = Path(folder) / "automotive-analysis.pptx"
            render_analysis(settings, analysis, plan, output_path)
            deck = Presentation(output_path)
            charts = [shape for slide in deck.slides for shape in slide.shapes if shape.has_chart]
            tables = [shape for slide in deck.slides for shape in slide.shapes if shape.has_table]
            pictures = [shape for slide in deck.slides for shape in slide.shapes if shape.shape_type == 13]
            self.assertEqual(len(deck.slides), len(plan["slide_ids"]))
            self.assertGreaterEqual(len(charts), 5)
            self.assertGreaterEqual(len(tables), 3)
            self.assertEqual(len(pictures), len(deck.slides))
            for slide in deck.slides:
                self.assertEqual(slide.background.fill.fore_color.rgb, PALE)
                logo = next(shape for shape in slide.shapes if shape.shape_type == 13)
                self.assertAlmostEqual(logo.left / Inches(1), 11.16, places=2)
                self.assertAlmostEqual(logo.top / Inches(1), 6.74, places=2)
                self.assertEqual(logo.image.blob, BRAND_LOGO.read_bytes())
            all_text = "\n".join(shape.text for slide in deck.slides for shape in slide.shapes
                                 if shape.has_text_frame)
            self.assertIn("AMO  /  CUSTOMER RESEARCH", all_text)
            first_slide_text = "\n".join(shape.text for shape in deck.slides[0].shapes if shape.has_text_frame)
            for slide in deck.slides:
                slide_text = "\n".join(shape.text for shape in slide.shapes if shape.has_text_frame)
                self.assertIn("Source: data-set.csv", slide_text)
                self.assertIn("Question answered:", slide_text)
                self.assertIn("TAKEAWAY", slide_text)
                source_footer = next(shape for shape in slide.shapes
                                     if shape.has_text_frame and "Source:" in shape.text)
                question_footer = next(shape for shape in slide.shapes
                                        if shape.has_text_frame and "Question answered:" in shape.text)
                self.assertLess(source_footer.top, question_footer.top)
                self.assertLess(source_footer.left, Inches(1))
                self.assertGreater(source_footer.top, Inches(6))
                takeaway_shape = next(shape for shape in slide.shapes
                                      if shape.has_text_frame and "TAKEAWAY" in shape.text)
                self.assertGreater(takeaway_shape.top, Inches(1.5))
                self.assertLess(takeaway_shape.top, Inches(2.1))
            self.assertIn("Respondent base", all_text)
            self.assertTrue(any(shape.text == "203" for shape in deck.slides[0].shapes if shape.has_text_frame))
            self.assertIn("51%", all_text)
            self.assertIn("5%", all_text)
            self.assertIn("demographics", all_text)
            self.assertIn("purchase intent/rejection", all_text)
            self.assertIn("Respondent base", first_slide_text)
            self.assertTrue(any(shape.text == "203" for shape in deck.slides[0].shapes if shape.has_text_frame))
            self.assertIn("51%", first_slide_text)
            self.assertEqual(sum(shape.has_table for shape in deck.slides[0].shapes), 1)
            self.assertTrue(any(shape.has_chart and shape.chart.chart_type == XL_CHART_TYPE.DOUGHNUT
                                for shape in deck.slides[0].shapes))
            self.assertTrue(any(shape.text == "Exterior design concept earns strong ratings" for shape in deck.slides[0].shapes
                                if shape.has_text_frame))
            fonts = {paragraph.font.name for slide in deck.slides for shape in slide.shapes
                     if shape.has_text_frame for paragraph in shape.text_frame.paragraphs
                     if paragraph.text.strip() and paragraph.font.name}
            fonts.update(run.font.name for slide in deck.slides for shape in slide.shapes
                         if shape.has_text_frame for paragraph in shape.text_frame.paragraphs
                         for run in paragraph.runs if run.text.strip() and run.font.name)
            self.assertIn(HEADING_FONT, fonts)
            self.assertIn(BODY_FONT, fonts)
            self.assertIn(
                "The design resonates with a majority: 51% rate it 9-10, rising to 63% among current BEV owners.",
                first_slide_text,
            )
            self.assertIn("BEV", " ".join(cell.text for table in tables for row in table.table.rows for cell in row.cells))
            with ZipFile(output_path) as archive:
                chart_path = next(name for name in archive.namelist()
                                  if name.startswith("ppt/charts/chart") and name.endswith(".xml"))
                root = etree.fromstring(archive.read(chart_path))
                namespace = "{http://schemas.openxmlformats.org/drawingml/2006/chart}"
                doughnut = root.find(f".//{namespace}doughnutChart")
                self.assertEqual(doughnut.find(f"{namespace}firstSliceAng").get("val"), "270")
                self.assertEqual(doughnut.find(f"{namespace}holeSize").get("val"), "72")

    def test_xlsx_without_vehicle_specific_headers_is_analyzed_dynamically(self):
        with tempfile.TemporaryDirectory() as folder:
            input_path = Path(folder) / "different.xlsx"
            output_path = Path(folder) / "result.pptx"
            workbook = Workbook()
            controls = workbook.active
            controls.title = "Presentation"
            controls.append(("Setting", "Value"))
            controls.append(("Data sheet", "Product feedback"))
            controls.append(("First slide", "Product review"))
            controls.append(("Title", "Autumn release"))
            data = workbook.create_sheet("Product feedback")
            data.append(("Customer segment", "Concept rating", "Vehicle interest"))
            for segment, rating, interest in (
                ("Commuter", 8, "SUV"), ("Family", 6, "SUV"),
                ("Commuter", 9, "Hatchback"), ("Family", 5, "SUV"),
                ("Commuter", 7, "Sedan"), ("Family", 4, "Hatchback"),
                ("Commuter", 8, "SUV"), ("Family", 7, "Sedan"),
                ("Commuter", 9, "Hatchback"), ("Family", 6, "SUV"),
                ("Commuter", 8, "Sedan"), ("Family", 5, "Hatchback"),
                ("Commuter", 7, "SUV"), ("Family", 6, "Sedan"),
                ("Commuter", 9, "Hatchback"), ("Family", 4, "SUV"),
                ("Commuter", 8, "Sedan"), ("Family", 7, "Hatchback"),
                ("Commuter", 7, "SUV"), ("Family", 5, "Sedan"),
            ):
                data.append((segment, rating, interest))
            workbook.save(input_path)

            settings, headers, records, labels = load_survey(input_path)
            analysis = analyze_survey(headers, records, labels)
            plan = validate_analysis_plan(default_analysis_plan(analysis), analysis)
            render_analysis(settings, analysis, plan, output_path)

            deck = Presentation(output_path)
            self.assertEqual(settings["Title"], "Autumn release")
            self.assertGreaterEqual(len(deck.slides), 2)
            self.assertGreaterEqual(sum(shape.has_chart for slide in deck.slides for shape in slide.shapes), 2)
            self.assertGreaterEqual(sum(shape.has_table for slide in deck.slides for shape in slide.shapes), 1)

    def test_unseen_one_header_csv_and_two_row_xlsx_keep_semantic_fields(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            csv_path = root / "new-market-study.csv"
            csv_lines = ["age;buyer_intent;price_expectation;Q_style"]
            for index in range(80):
                csv_lines.append(f"{25 + index % 30};{'Yes' if index % 3 else 'No'};"
                                 f"{25000 + index % 4 * 5000};{1 + index % 10}")
            csv_path.write_text("\n".join(csv_lines), encoding="utf-8")
            _, csv_fields, csv_records, csv_labels = load_survey(csv_path)
            csv_analysis = analyze_survey(csv_fields, csv_records, csv_labels)
            self.assertIn("rating:Q_style", csv_analysis["slides"])
            self.assertIn("frequency:buyer_intent", csv_analysis["slides"])
            self.assertIn("numeric:age", csv_analysis["slides"])
            self.assertIn("numeric:price_expectation", csv_analysis["slides"])
            self.assertNotIn("rating:buyer_intent", csv_analysis["slides"])
            self.assertTrue(csv_analysis["coverage"]["demographics"])
            self.assertTrue(csv_analysis["coverage"]["purchase_intent_or_rejection"])
            self.assertTrue(csv_analysis["coverage"]["pricing_expectations"])

            xlsx_path = root / "two-header-survey.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Responses"
            sheet.append(("age", "buyer_intent", "Q_style"))
            sheet.append(("Age of respondent", "Would you consider buying?",
                          "How do you rate the concept from 1 to 10?"))
            for index in range(40):
                sheet.append((25 + index % 20, "Yes" if index % 3 else "No", 1 + index % 10))
            workbook.save(xlsx_path)
            _, xlsx_fields, xlsx_records, xlsx_labels = load_survey(xlsx_path)
            self.assertEqual(xlsx_labels["age"], "Age of respondent")
            self.assertEqual(xlsx_labels["Q_style"], "How do you rate the concept from 1 to 10?")
            self.assertEqual(len(xlsx_records), 40)

    def test_plan_rejects_uncomputed_slide_ids(self):
        with self.assertRaisesRegex(ValueError, "unavailable computed slide"):
            validate_analysis_plan({"slide_ids": ["rating:made-up"], "insight_ids": []},
                                   {"slides": {}, "required_ratings": [], "insights": {}},
                                   strict=True)
        # Without strict mode, gracefully filters out unavailable IDs and uses safe fallback
        resilient = validate_analysis_plan(
            {"slide_ids": ["rating:made-up"], "insight_ids": ["non-existent"]},
            {"slides": {"frequency:a": {"kind": "frequency"}},
             "required_ratings": [], "insights": {"insight:1": "test"}}
        )
        self.assertIn("frequency:a", resilient["slide_ids"])

    def test_numeric_demographics_and_purchase_rejection_cross_tabs(self):
        records = []
        for index in range(40):
            records.append({"Age": 20 + index, "Gender": "Woman" if index % 2 else "Man",
                            "Would buy": "Yes" if index % 4 else "No",
                            "Styling rating": 6 + index % 5})
        labels = {"Age": "Age of respondent", "Gender": "Gender",
                  "Would buy": "Would you buy this vehicle?",
                  "Styling rating": "Styling rating on a 10-point scale"}
        analysis = analyze_survey(list(labels), records, labels)
        self.assertTrue(analysis["coverage"]["demographics"])
        self.assertTrue(analysis["coverage"]["purchase_intent_or_rejection"])
        self.assertIn("numeric:Age", analysis["slides"])
        age_crosstab = analysis["slides"]["crosstab:Would buy:Age"]
        self.assertGreaterEqual(len(age_crosstab["groups"]), 2)
        plan = default_analysis_plan(analysis)
        with tempfile.TemporaryDirectory() as folder:
            output = Path(folder) / "buyer-profile.pptx"
            render_analysis({"Data sheet": "Customer survey"}, analysis, plan, output)
            deck = Presentation(output)
            self.assertGreaterEqual(sum(shape.has_chart for slide in deck.slides for shape in slide.shapes), 3)

    def test_empty_widescreen_template_renders_native_chart(self):
        with tempfile.TemporaryDirectory() as folder:
            template = Path(folder) / "template.pptx"
            output = Path(folder) / "out.pptx"
            source = Presentation()
            source.slide_width = Inches(13.333)
            source.slide_height = Inches(7.5)
            source.save(template)
            settings = {"First slide": "Overview", "Title": "Survey", "Audience": "Team",
                        "Data sheet": "Answers"}
            records = [{"Segment": "A", "Rating": 8}, {"Segment": "B", "Rating": 9},
                       {"Segment": "A", "Rating": 7}, {"Segment": "B", "Rating": 10},
                       {"Segment": "A", "Rating": 9}, {"Segment": "B", "Rating": 8},
                       {"Segment": "A", "Rating": 8}, {"Segment": "B", "Rating": 9},
                       {"Segment": "A", "Rating": 7}, {"Segment": "B", "Rating": 10},
                       {"Segment": "A", "Rating": 9}, {"Segment": "B", "Rating": 8},
                       {"Segment": "A", "Rating": 8}, {"Segment": "B", "Rating": 9},
                       {"Segment": "A", "Rating": 7}, {"Segment": "B", "Rating": 10},
                       {"Segment": "A", "Rating": 9}, {"Segment": "B", "Rating": 8},
                       {"Segment": "A", "Rating": 8}, {"Segment": "B", "Rating": 9}]
            analysis = analyze_survey(["Segment", "Rating"], records,
                                      {"Segment": "Customer segment", "Rating": "Purchase rating"})
            plan = default_analysis_plan(analysis)
            render_analysis(settings, analysis, plan, output, template)
            deck = Presentation(output)
            self.assertGreaterEqual(len(deck.slides), 2)
            self.assertGreaterEqual(sum(shape.has_chart for slide in deck.slides for shape in slide.shapes), 2)


if __name__ == "__main__":
    unittest.main()