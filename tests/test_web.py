import os
import tempfile
import time
import unittest
from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import Workbook
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches

from survey_studio import app as web


class WebWorkflowTests(unittest.TestCase):
    def setUp(self):
        self.storage = tempfile.TemporaryDirectory()
        self.old_uploads, self.old_generated = web.UPLOADS, web.GENERATED
        web.UPLOADS = Path(self.storage.name) / "uploads"
        web.GENERATED = Path(self.storage.name) / "generated"
        self.client = TestClient(web.app)
        self.client.__enter__()

    def tearDown(self):
        self.client.__exit__(None, None, None)
        web.UPLOADS, web.GENERATED = self.old_uploads, self.old_generated
        self.storage.cleanup()

    def workbook(self):
        source = Path(self.storage.name) / "input.xlsx"
        book = Workbook()
        first = book.active
        first.title = "First"
        first.append(("Score",))
        first.append(("Low",))
        second = book.create_sheet("Second")
        second.append(("Response", "Comment"))
        for answer in ("Yes", "No", "Yes", "No"):
            second.append((answer, "This is a full survey comment about the experience."))
        controls = book.create_sheet("Presentation")
        controls.append(("Setting", "Value"))
        controls.append(("Output name", "customer-brief.PPTX"))
        controls.append(("Data sheet", "Second"))
        book.save(source)
        return source

    def upload(self, route, source, data=None):
        with source.open("rb") as stream:
            return self.client.post(route, files={"file": ("test.xlsx", stream)}, data=data or {})

    def test_select_generate_download_and_remove_upload(self):
        source = self.workbook()
        inspection = self.upload("/api/inspect", source)
        self.assertEqual(inspection.status_code, 200)
        self.assertEqual(inspection.json(), {"sheets": ["First", "Second"], "suggested": "Second"})
        self.assertEqual(list(web.UPLOADS.iterdir()), [])
        generated = self.upload("/api/generate", source, {"sheet": "Second", "ai": "false"})
        self.assertEqual(generated.status_code, 200, generated.text)
        self.assertEqual(generated.json()["filename"], "customer-brief.pptx")
        self.assertEqual(list(web.UPLOADS.iterdir()), [])
        download = self.client.get(generated.json()["download"])
        self.assertEqual(download.status_code, 200)
        self.assertEqual(len(Presentation(web.GENERATED.glob("*.pptx").__next__()).slides), 2)
        self.assertIn("customer-brief.pptx", download.headers["content-disposition"])

    def test_bad_selection_does_not_store_workbook(self):
        source = self.workbook()
        response = self.upload("/api/generate", source, {"sheet": "Missing"})
        self.assertEqual(response.status_code, 400)
        self.assertEqual(list(web.UPLOADS.iterdir()), [])
        self.assertEqual(list(web.GENERATED.iterdir()), [])

    def test_invalid_excel_file_returns_client_error_and_is_removed(self):
        response = self.client.post("/api/inspect", files={"file": ("fake.xlsx", b"not a zip file")})
        self.assertEqual(response.status_code, 400)
        self.assertEqual(list(web.UPLOADS.iterdir()), [])

    def test_expired_decks_are_removed_and_not_downloadable(self):
        web.GENERATED.mkdir(parents=True, exist_ok=True)
        name = "a" * 32
        old = web.GENERATED / f"{name}--old.pptx"
        old.write_bytes(b"old")
        timestamp = time.time() - web.RETENTION_SECONDS - 10
        os.utime(old, (timestamp, timestamp))
        self.assertEqual(self.client.get(f"/api/files/{name}").status_code, 404)
        self.assertFalse(old.exists())
        stale = web.GENERATED / f"{'b' * 32}--stale.pptx"
        stale.write_bytes(b"old")
        os.utime(stale, (timestamp, timestamp))
        web.cleanup_expired()
        self.assertFalse(stale.exists())

    def test_available_template_preview_and_generation(self):
        listing = self.client.get("/api/templates")
        self.assertEqual(listing.status_code, 200)
        template_id = listing.json()[0]["id"]
        preview = self.client.get(f"/api/templates/{template_id}/preview")
        self.assertEqual(preview.status_code, 200)
        self.assertEqual(len(preview.json()["slides"]), 2)
        generated = self.upload("/api/generate", self.workbook(),
                                {"sheet": "Second", "ai": "false", "template_id": template_id})
        self.assertEqual(generated.status_code, 200, generated.text)
        deck = Presentation(next(web.GENERATED.glob("*.pptx")))
        self.assertEqual(len(deck.slides), 2)
        self.assertEqual(deck.slides[0].background.fill.fore_color.rgb, RGBColor(255, 255, 255))

    def test_uploaded_multislide_template_preview_and_generation(self):
        template = Path(self.storage.name) / "own.pptx"
        presentation = Presentation()
        presentation.slide_width = Inches(13.333)
        presentation.slide_height = Inches(7.5)
        for color in ("EDF5FA", "FFFFFF", "F1F7FA"):
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            slide.background.fill.solid()
            slide.background.fill.fore_color.rgb = RGBColor.from_string(color)
        presentation.save(template)
        with template.open("rb") as stream:
            response = self.client.post("/api/templates/preview", files={"file": ("own.pptx", stream)})
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(len(response.json()["slides"]), 3)
        self.assertEqual(list(web.UPLOADS.iterdir()), [])
        with template.open("rb") as stream, self.workbook().open("rb") as workbook:
            generated = self.client.post("/api/generate", data={"sheet": "Second", "ai": "false"},
                                         files={"file": ("data.xlsx", workbook), "template_file": ("own.pptx", stream)})
        self.assertEqual(generated.status_code, 200, generated.text)
        self.assertEqual(list(web.UPLOADS.iterdir()), [])
        deck = Presentation(next(web.GENERATED.glob("*.pptx")))
        self.assertEqual(len(deck.slides), 2)
        self.assertEqual(deck.slides[0].background.fill.fore_color.rgb, RGBColor.from_string("EDF5FA"))

    def test_invalid_template_does_not_leave_uploads(self):
        response = self.client.post("/api/templates/preview", files={"file": ("bad.pptx", b"not a pptx")})
        self.assertEqual(response.status_code, 400)
        self.assertEqual(list(web.UPLOADS.iterdir()), [])


if __name__ == "__main__":
    unittest.main()