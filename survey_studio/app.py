"""Local survey-to-PowerPoint web application."""

import asyncio
import base64
import os
import re
import time
from contextlib import asynccontextmanager
from pathlib import Path
from uuid import uuid4
from zipfile import BadZipFile

from lxml import etree
from fastapi import FastAPI, File, Form, HTTPException, UploadFile
from fastapi.responses import FileResponse
from fastapi.staticfiles import StaticFiles
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
from .smart_slides import (BODY_FONT, BRAND_LOGO, HEADING_FONT, INK, LIGHT_BLUE, PALE, TEAL)
from .smart_slides import (ai_analysis_plan, analyze_survey, default_analysis_plan,
                           load_survey, render_analysis, validate_analysis_plan)

STORAGE = Path(os.environ.get("PPT_STORAGE_DIR", ROOT / "storage"))
UPLOADS = STORAGE / "uploads"
GENERATED = STORAGE / "generated"
TEMPLATES = STORAGE / "templates"
AVAILABLE_TEMPLATES = {"hkm-light": "HKM light", "hkm-grey": "HKM cool grey"}
MAX_UPLOAD = 15 * 1024 * 1024
RETENTION_SECONDS = 7 * 24 * 60 * 60
FILE_ID = re.compile(r"^[0-9a-f]{32}$")


def cleanup_expired(now=None):
    now = time.time() if now is None else now
    for directory, lifetime in ((UPLOADS, 24 * 60 * 60), (GENERATED, RETENTION_SECONDS)):
        directory.mkdir(parents=True, exist_ok=True)
        for path in directory.iterdir():
            if path.is_file() and now - path.stat().st_mtime >= lifetime:
                path.unlink(missing_ok=True)


def _apply_hkm_theme(presentation):
    theme_part = next(
        relation.target_part for relation in presentation.slide_master.part.rels.values()
        if relation.reltype.endswith("/theme")
    )
    root = etree.fromstring(theme_part.blob)
    namespaces = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    root.find(".//a:fontScheme/a:majorFont/a:latin", namespaces).set("typeface", HEADING_FONT)
    root.find(".//a:fontScheme/a:minorFont/a:latin", namespaces).set("typeface", BODY_FONT)

    colors = {"dk1": "595959", "lt1": "F2F2F2", "accent1": "013B67",
              "accent2": "DBE3ED", "accent3": "595959"}
    for slot, color in colors.items():
        slot_element = root.find(f".//a:clrScheme/a:{slot}", namespaces)
        color_element = next(iter(slot_element))
        color_element.tag = f"{{{namespaces['a']}}}srgbClr"
        color_element.attrib.clear()
        color_element.set("val", color)
    theme_part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def create_available_templates():
    TEMPLATES.mkdir(parents=True, exist_ok=True)
    for template_id, colors in {
        "hkm-light": ("F2F2F2", "F2F2F2"),
        "hkm-grey": ("D3D9E6", "F2F2F2"),
    }.items():
        path = TEMPLATES / f"{template_id}.pptx"
        presentation = Presentation()
        presentation.slide_width = Inches(13.333)
        presentation.slide_height = Inches(7.5)
        for number, color in enumerate(colors, start=1):
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            slide.background.fill.solid()
            slide.background.fill.fore_color.rgb = RGBColor.from_string(color)
            if BRAND_LOGO.is_file():
                slide.shapes.add_picture(str(BRAND_LOGO), Inches(11.16), Inches(6.74),
                                         width=Inches(1.8), height=Inches(0.62))
        _apply_hkm_theme(presentation)
        presentation.save(path)


def template_preview(path):
    presentation = Presentation(path)
    if abs(presentation.slide_width / presentation.slide_height - 16 / 9) > 0.02:
        raise ValueError("Template must use a 16:9 slide size.")
    if len(presentation.slides) > 30:
        raise ValueError("Template has too many slides (maximum 30).")
    slides = []
    for slide in presentation.slides:
        background = "#ffffff"
        try:
            if slide.background.fill.type is not None and slide.background.fill.fore_color.rgb:
                background = f"#{slide.background.fill.fore_color.rgb}"
        except (AttributeError, TypeError):
            pass
        items = []
        for shape in slide.shapes:
            if len(items) >= 35:
                break
            item = {"x": max(0, shape.left / presentation.slide_width * 100),
                    "y": max(0, shape.top / presentation.slide_height * 100),
                    "w": max(0, shape.width / presentation.slide_width * 100),
                    "h": max(0, shape.height / presentation.slide_height * 100)}
            if shape.has_text_frame and shape.text.strip():
                run = next((run for paragraph in shape.text_frame.paragraphs for run in paragraph.runs
                            if run.text.strip()), None)
                font_color = "#595959"
                if run:
                    try:
                        if run.font.color.rgb:
                            font_color = f"#{run.font.color.rgb}"
                    except (AttributeError, TypeError, ValueError):
                        pass
                item.update(kind="text", text=shape.text[:150],
                            font=run.font.name if run and run.font.name else BODY_FONT,
                            color=font_color)
            elif shape.shape_type == 13 and shape.image.content_type.startswith("image/"):
                item.update(kind="image", data=f"data:{shape.image.content_type};base64,"
                            f"{base64.b64encode(shape.image.blob).decode('ascii')}")
            else:
                try:
                    color = shape.fill.fore_color.rgb
                except (AttributeError, TypeError):
                    color = None
                if color is None:
                    continue
                item.update(kind="block", color=f"#{color}")
            items.append(item)
        slides.append({"background": background, "items": items})
    return {"slides": slides}


@asynccontextmanager
async def lifespan(app):
    cleanup_expired()
    create_available_templates()
    stop = asyncio.Event()

    async def sweeper():
        while not stop.is_set():
            try:
                await asyncio.wait_for(stop.wait(), timeout=3600)
            except asyncio.TimeoutError:
                cleanup_expired()

    task = asyncio.create_task(sweeper())
    try:
        yield
    finally:
        stop.set()
        await task


app = FastAPI(title="HKM Survey Studio", lifespan=lifespan)
app.mount("/static", StaticFiles(directory=Path(__file__).parent / "static"), name="static")


async def save_upload(file, destination, extension=".xlsx"):
    if not file.filename or Path(file.filename).suffix.lower() != extension:
        raise HTTPException(400, f"Select a {extension} file.")
    total = 0
    try:
        with destination.open("wb") as stream:
            while chunk := await file.read(1024 * 1024):
                total += len(chunk)
                if total > MAX_UPLOAD:
                    raise HTTPException(413, "Workbook exceeds the 15 MB limit.")
                stream.write(chunk)
    finally:
        await file.close()
    if not total:
        raise HTTPException(400, "Workbook is empty.")


def sheet_names(path):
    if Path(path).suffix.lower() == ".csv":
        return ["Data"]
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheets = [name for name in workbook.sheetnames if name != "Presentation"]
        if not sheets:
            raise ValueError("No data worksheets found in this workbook.")
        return sheets
    finally:
        workbook.close()


def suggested_sheet(path, names):
    if Path(path).suffix.lower() == ".csv":
        return "Data"
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if "Presentation" in workbook:
            for key, value in workbook["Presentation"].iter_rows(min_row=2, max_col=2, values_only=True):
                if key == "Data sheet" and value in names:
                    return value
        return names[0]
    finally:
        workbook.close()


def plan_deck(path, sheet, ai, source_name=None, model=None):
    model = model or os.environ.get("OLLAMA_MODEL", "llama3.1:8b")
    if sheet not in sheet_names(path):
        raise ValueError("Selected worksheet does not exist.")
    settings, headers, records, labels = load_survey(path, sheet_name=sheet)
    if source_name:
        settings["Source name"] = source_name
    analysis = analyze_survey(headers, records, labels)
    plan = validate_analysis_plan(ai_analysis_plan(analysis, settings, model) if ai
                                  else default_analysis_plan(analysis), analysis)
    return settings, analysis, plan


@app.get("/")
def home():
    return FileResponse(Path(__file__).parent / "static" / "index.html")


@app.get("/api/templates")
def available_templates():
    create_available_templates()
    return [{"id": template_id, "name": name} for template_id, name in AVAILABLE_TEMPLATES.items()]


@app.get("/api/templates/{template_id}/preview")
def available_preview(template_id: str):
    if template_id not in AVAILABLE_TEMPLATES:
        raise HTTPException(404, "Template not found.")
    return template_preview(TEMPLATES / f"{template_id}.pptx")


@app.post("/api/templates/preview")
async def uploaded_preview(file: UploadFile = File(...)):
    UPLOADS.mkdir(parents=True, exist_ok=True)
    temporary = UPLOADS / f"{uuid4().hex}.pptx"
    try:
        await save_upload(file, temporary, ".pptx")
        return await asyncio.to_thread(template_preview, temporary)
    except (ValueError, OSError, BadZipFile) as error:
        raise HTTPException(400, f"Cannot preview template: {error}") from error
    finally:
        temporary.unlink(missing_ok=True)


@app.post("/api/inspect")
async def inspect(file: UploadFile = File(...)):
    UPLOADS.mkdir(parents=True, exist_ok=True)
    extension = Path(file.filename or "").suffix.lower()
    if extension not in (".xlsx", ".csv"):
        await file.close()
        raise HTTPException(400, "Select an .xlsx or .csv survey file.")
    temporary = UPLOADS / f"{uuid4().hex}{extension}"
    try:
        await save_upload(file, temporary, extension)
        names = await asyncio.to_thread(sheet_names, temporary)
        return {"sheets": names, "suggested": await asyncio.to_thread(suggested_sheet, temporary, names)}
    except (ValueError, OSError, BadZipFile, InvalidFileException) as error:
        raise HTTPException(400, f"Cannot read workbook: {error}") from error
    finally:
        temporary.unlink(missing_ok=True)


@app.post("/api/generate")
async def generate(file: UploadFile = File(...), sheet: str = Form(...), ai: bool = Form(False),
                   template_id: str = Form(""), template_file: UploadFile | None = File(None)):
    UPLOADS.mkdir(parents=True, exist_ok=True)
    GENERATED.mkdir(parents=True, exist_ok=True)
    extension = Path(file.filename or "").suffix.lower()
    if extension not in (".xlsx", ".csv"):
        await file.close()
        raise HTTPException(400, "Select an .xlsx or .csv survey file.")
    temporary = UPLOADS / f"{uuid4().hex}{extension}"
    temporary_template = UPLOADS / f"{uuid4().hex}.pptx"
    output = None
    try:
        await save_upload(file, temporary, extension)
        if template_file and template_id:
            raise ValueError("Choose either an uploaded or an available template.")
        template = None
        if template_file:
            await save_upload(template_file, temporary_template, ".pptx")
            await asyncio.to_thread(template_preview, temporary_template)
            template = temporary_template
        elif template_id:
            if template_id not in AVAILABLE_TEMPLATES:
                raise ValueError("Selected template is not available.")
            template = TEMPLATES / f"{template_id}.pptx"
        settings, analysis, plan = await asyncio.to_thread(plan_deck, temporary, sheet, ai, file.filename)
        name = settings["Output name"]
        if extension == ".csv":
            name = f"{Path(file.filename).stem}-presentation.pptx"
        if not name or Path(name).name != name or Path(name).suffix.lower() != ".pptx":
            raise ValueError("Output name in Presentation must be a .pptx filename only.")
        download_name = re.sub(r"[^\w. -]", "_", Path(name).stem).strip(" .") + ".pptx"
        job_id = uuid4().hex
        output = GENERATED / f"{job_id}--{download_name}"
        await asyncio.to_thread(render_analysis, settings, analysis, plan, output, template)
        return {"id": job_id, "filename": download_name,
            "respondents": analysis["n"],
            "slides": len(plan["slide_ids"]),
            "charts": len(plan["slide_ids"]), "download": f"/api/files/{job_id}"}
    except (ValueError, OSError, BadZipFile, InvalidFileException) as error:
        if output:
            output.unlink(missing_ok=True)
        raise HTTPException(400, str(error)) from error
    finally:
        temporary.unlink(missing_ok=True)
        temporary_template.unlink(missing_ok=True)


@app.get("/api/files/{job_id}")
def download(job_id: str):
    if not FILE_ID.fullmatch(job_id):
        raise HTTPException(404, "Presentation not found.")
    matches = list(GENERATED.glob(f"{job_id}--*.pptx"))
    if not matches or time.time() - matches[0].stat().st_mtime >= RETENTION_SECONDS:
        for path in matches:
            path.unlink(missing_ok=True)
        raise HTTPException(404, "Presentation expired or not found.")
    return FileResponse(matches[0], filename=matches[0].name.split("--", 1)[1],
                        media_type="application/vnd.openxmlformats-officedocument.presentationml.presentation")